"""Editable CV model: parse extracted text into sections, add a skill safely, export DOCX / PDF.

The original uploaded CV is never modified.  When the original is a PDF, the CV is rebuilt from the
extracted text, so the original PDF layout is NOT preserved.
"""
from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field
from html import escape

from core.skills import extract_skills, normalize_skill

KNOWN_HEADINGS = {
    "summary", "professional summary", "career summary", "objective", "career objective", "profile",
    "professional profile", "about me", "about", "education", "academic background", "academic qualifications",
    "experience", "work experience", "professional experience", "employment history", "work history",
    "internships", "internship", "projects", "personal projects", "academic projects", "key projects",
    "skills", "technical skills", "core skills", "key skills", "core competencies", "skills summary",
    "certifications", "certificates", "licenses", "courses", "training", "achievements", "awards",
    "honors", "publications", "languages", "interests", "hobbies", "references", "contact",
    "contact information", "personal details", "personal information", "volunteer experience",
    "extracurricular activities", "activities",
}
SKILL_SECTION_HINTS = ("skill", "competenc")
BULLETS = ("•", "-", "*", "–", "▪", "●", "◦")


@dataclass
class Section:
    heading: str | None
    lines: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {"heading": self.heading, "lines": list(self.lines)}


def is_heading(line: str) -> bool:
    stripped = line.strip().rstrip(":").strip()
    if not stripped or len(stripped) > 45:
        return False
    return stripped.lower() in KNOWN_HEADINGS


def parse_cv(text: str) -> list[Section]:
    sections = [Section(None, [])]
    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if is_heading(line):
            sections.append(Section(line.strip().rstrip(":").strip(), []))
        else:
            sections[-1].lines.append(line)
    for sec in sections:                                   # trim blank edges
        while sec.lines and not sec.lines[0].strip():
            sec.lines.pop(0)
        while sec.lines and not sec.lines[-1].strip():
            sec.lines.pop()
    return [s for s in sections if s.heading or s.lines]


def sections_from_dicts(items: list[dict]) -> list[Section]:
    return [Section(i.get("heading"), list(i.get("lines", []))) for i in items]


def sections_to_dicts(sections: list[Section]) -> list[dict]:
    return [s.to_dict() for s in sections]


def cv_to_text(sections: list[Section]) -> str:
    blocks = []
    for sec in sections:
        block = ([sec.heading] if sec.heading else []) + sec.lines
        blocks.append("\n".join(x for x in block if x is not None).strip("\n"))
    return "\n\n".join(b for b in blocks if b).strip() + "\n"


def section_names(sections: list[Section]) -> list[str]:
    return [s.heading for s in sections if s.heading]


def skills_section_options(sections: list[Section]) -> list[str]:
    """Existing skill-like sections first, then the common defaults that do not exist yet."""
    names = section_names(sections)
    existing = [n for n in names if any(h in n.lower() for h in SKILL_SECTION_HINTS)]
    defaults = [d for d in ("Technical Skills", "Skills") if d.lower() not in {n.lower() for n in names}]
    others = [n for n in names if n not in existing]
    return existing + defaults + others


def _find(sections: list[Section], name: str) -> Section | None:
    for sec in sections:
        if sec.heading and sec.heading.lower() == name.lower():
            return sec
    return None


def skill_already_in_cv(sections: list[Section], skill_text: str) -> bool:
    """True when the CV already mentions the skill (by canonical name or by exact wording)."""
    text = cv_to_text(sections)
    canonical = normalize_skill(skill_text)
    if canonical in extract_skills(text):
        return True
    return bool(re.search(r"(?<![A-Za-z0-9])" + re.escape(skill_text.strip()) + r"(?![A-Za-z0-9])", text, re.I))


def _looks_like_list(line: str) -> bool:
    s = line.strip()
    return "," in s and not s.endswith(".") and len(s.split()) <= 45 and not s.startswith(BULLETS)


def add_skill(sections: list[Section], skill_text: str, section_name: str) -> list[Section]:
    """Return a NEW list of sections with `skill_text` added; unrelated content is untouched."""
    skill_text = skill_text.strip()
    if not skill_text:
        raise ValueError("The skill wording cannot be empty.")
    new = copy.deepcopy(sections)
    target = _find(new, section_name)
    if target is None:
        new.append(Section(section_name.strip(), [skill_text]))
        return new
    idx = max((i for i, ln in enumerate(target.lines) if ln.strip()), default=None)
    if idx is None:
        target.lines.append(skill_text)
        return new
    last = target.lines[idx]
    stripped = last.lstrip()
    if stripped.startswith(BULLETS):
        indent = last[: len(last) - len(stripped)]
        target.lines.insert(idx + 1, f"{indent}{stripped[0]} {skill_text}")
    elif _looks_like_list(last):
        target.lines[idx] = last.rstrip(" ,;") + f", {skill_text}"
    else:
        target.lines.insert(idx + 1, skill_text)
    return new


def section_preview(sections: list[Section], name: str) -> str:
    sec = _find(sections, name)
    if sec is None:
        return "(this section does not exist yet and will be created)"
    return "\n".join([sec.heading] + sec.lines).strip()


# ---- export -----------------------------------------------------------------------------------------
def _split_header(sections: list[Section]):
    if sections and sections[0].heading is None:
        lines = [ln for ln in sections[0].lines if ln.strip()]
        return (lines[0] if lines else "", lines[1:], sections[1:])
    return ("", [], sections)


def _is_bullet(line: str) -> bool:
    return line.lstrip().startswith(BULLETS)


def _unbullet(line: str) -> str:
    return line.lstrip()[1:].strip()


def to_docx(sections: list[Section]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.75)
        sec.top_margin = sec.bottom_margin = Inches(0.65)

    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10)
    navy = RGBColor(0x0F, 0x17, 0x2A)       # Slate 900
    indigo = RGBColor(0x1E, 0x3A, 0x8A)     # Deep Corporate Navy
    muted = RGBColor(0x47, 0x55, 0x69)      # Slate 600

    def add_bottom_border(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'CBD5E1')
        pBdr.append(bottom)
        pPr.append(pBdr)

    name, contact, rest = _split_header(sections)
    if name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(name)
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = navy

    if contact:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        joined_contact = "  •  ".join(ln.strip() for ln in contact if ln.strip())
        run = p.add_run(joined_contact)
        run.font.size = Pt(9.5)
        run.font.color.rgb = muted
        add_bottom_border(p)

    for sec in rest:
        if sec.heading:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(sec.heading.upper())
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = indigo
            add_bottom_border(p)

        for line in sec.lines:
            if not line.strip():
                continue
            if _is_bullet(line):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.left_indent = Inches(0.25)
                text = _unbullet(line)
                # If there's a bold prefix like "Languages: Python, Go"
                if ":" in text and len(text.split(":", 1)[0].split()) <= 4:
                    prefix, rest_text = text.split(":", 1)
                    r_pre = p.add_run(prefix + ":")
                    r_pre.bold = True
                    r_pre.font.color.rgb = navy
                    r_rest = p.add_run(rest_text)
                    r_rest.font.color.rgb = navy
                else:
                    r = p.add_run(text)
                    r.font.color.rgb = navy
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2.5)
                text = line.strip()
                if ":" in text and len(text.split(":", 1)[0].split()) <= 4:
                    prefix, rest_text = text.split(":", 1)
                    r_pre = p.add_run(prefix + ":")
                    r_pre.bold = True
                    r_pre.font.color.rgb = navy
                    r_rest = p.add_run(rest_text)
                    r_rest.font.color.rgb = navy
                else:
                    r = p.add_run(text)
                    r.font.color.rgb = navy

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def to_pdf(sections: list[Section]) -> bytes:
    from reportlab.graphics.shapes import Drawing, Line
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    navy = colors.HexColor("#0F172A")
    indigo = colors.HexColor("#1E3A8A")
    muted = colors.HexColor("#475569")
    line_col = colors.HexColor("#CBD5E1")

    title_style = ParagraphStyle("t", parent=styles["Title"], textColor=navy, fontSize=20, leading=24,
                                 alignment=TA_CENTER, spaceAfter=2)
    contact_style = ParagraphStyle("c", parent=styles["BodyText"], alignment=TA_CENTER, fontSize=9, leading=12,
                                   textColor=muted, spaceAfter=6)
    head_style = ParagraphStyle("h", parent=styles["Heading3"], textColor=indigo, fontSize=11, leading=14,
                                spaceBefore=8, spaceAfter=2, fontName="Helvetica-Bold")
    body_style = ParagraphStyle("b", parent=styles["BodyText"], fontSize=9.5, leading=13.5, textColor=navy,
                                spaceBefore=0, spaceAfter=2.5)
    bullet_style = ParagraphStyle("bl", parent=body_style, leftIndent=14, bulletIndent=4, spaceBefore=0, spaceAfter=2)

    buffer = io.BytesIO()
    # 38pt margins for balanced page distribution
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=38, rightMargin=38, topMargin=36, bottomMargin=36)
    name, contact, rest = _split_header(sections)
    story = []

    if name:
        story.append(Paragraph(escape(name), title_style))
    if contact:
        joined_contact = "  &bull;  ".join(escape(ln.strip()) for ln in contact if ln.strip())
        story.append(Paragraph(joined_contact, contact_style))
        d = Drawing(doc.width, 2)
        d.add(Line(0, 0, doc.width, 0, strokeColor=line_col, strokeWidth=0.8))
        story.append(d)
        story.append(Spacer(1, 4))

    for sec in rest:
        if sec.heading:
            story.append(Paragraph(escape(sec.heading.upper()), head_style))
            d = Drawing(doc.width, 2)
            d.add(Line(0, 0, doc.width, 0, strokeColor=line_col, strokeWidth=0.8))
            story.append(d)
            story.append(Spacer(1, 2))

        for line in sec.lines:
            if not line.strip():
                story.append(Spacer(1, 3))
            elif _is_bullet(line):
                text = escape(_unbullet(line))
                if ":" in text and len(text.split(":", 1)[0].split()) <= 4:
                    prefix, rest_t = text.split(":", 1)
                    text = f"<b>{prefix}:</b>{rest_t}"
                story.append(Paragraph(text, bullet_style, bulletText="&bull;"))
            else:
                text = escape(line.strip())
                if ":" in text and len(text.split(":", 1)[0].split()) <= 4:
                    prefix, rest_t = text.split(":", 1)
                    text = f"<b>{prefix}:</b>{rest_t}"
                story.append(Paragraph(text, body_style))

    if not story:
        story.append(Paragraph("(empty CV)", body_style))
    doc.build(story)
    return buffer.getvalue()


# ---- in-place DOCX editing (keeps the original design) ---------------------------------------------
def _para_is_heading(paragraph) -> bool:
    text = paragraph.text.strip().rstrip(":").strip()
    if not text or len(text) > 45:
        return False
    if text.lower() in KNOWN_HEADINGS:
        return True
    return paragraph.style.name.lower().startswith("heading") or any(r.bold for r in paragraph.runs if r.text.strip()) and text.isupper()


def add_skill_to_docx(original: bytes, skill_text: str, section_name: str) -> bytes:
    """Return a copy of the ORIGINAL .docx with `skill_text` added to the skills section.

    * If a section with the given heading exists, the skill is appended to its last non-empty paragraph
      (comma list) or added as a new bullet/paragraph copying that paragraph's formatting.
    * If it does not exist, a heading + paragraph is inserted before the first Experience/Projects/Education
      section, or at the end when none is found.
    Everything else in the document is left untouched.
    """
    import copy

    from docx import Document

    skill_text = skill_text.strip()
    if not skill_text:
        raise ValueError("The skill wording cannot be empty.")
    doc = Document(io.BytesIO(original))
    paras = doc.paragraphs
    wanted = section_name.strip().lower()

    heading_idx = None
    for i, p in enumerate(paras):
        if _para_is_heading(p) and p.text.strip().rstrip(":").strip().lower() == wanted:
            heading_idx = i
            break

    if heading_idx is not None:
        # body paragraphs of that section: until the next heading
        body = []
        for p in paras[heading_idx + 1:]:
            if _para_is_heading(p):
                break
            if p.text.strip():
                body.append(p)
        if body:
            last = body[-1]
            if last.text.rstrip().endswith((",", ";")) or "," in last.text and not last.text.rstrip().endswith("."):
                _append_run(last, f", {skill_text}" if not last.text.rstrip().endswith((",", ";")) else f" {skill_text}")
            else:
                _insert_paragraph_after(last, skill_text, template=last)
        else:
            _insert_paragraph_after(paras[heading_idx], skill_text, template=paras[heading_idx])
    else:
        anchor = None
        for p in paras:
            if _para_is_heading(p) and p.text.strip().lower().rstrip(":") in {
                    "experience", "work experience", "professional experience", "projects", "education",
                    "employment history", "internships", "academic projects"}:
                anchor = p
                break
        heading_src = next((p for p in paras if _para_is_heading(p)), None) or (paras[0] if paras else None)
        body_src = next((p for p in paras if p.text.strip() and not _para_is_heading(p)), None) or heading_src
        if anchor is not None:
            new_heading = _insert_paragraph_before(anchor, section_name.strip(), template=heading_src)
            _insert_paragraph_after(new_heading, skill_text, template=body_src)
        else:
            new_heading = doc.add_paragraph(section_name.strip(), style=heading_src.style if heading_src else None)
            if heading_src is not None and heading_src.runs:
                _copy_run_format(heading_src.runs[0], new_heading.runs[0]) if new_heading.runs else None
            body = doc.add_paragraph(skill_text, style=body_src.style if body_src else None)
            if body_src is not None:
                body.paragraph_format.space_after = body_src.paragraph_format.space_after
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _copy_run_format(src_run, dst_run) -> None:
    dst_run.bold, dst_run.italic = src_run.bold, src_run.italic
    dst_run.font.size = src_run.font.size
    dst_run.font.name = src_run.font.name
    if src_run.font.color is not None and src_run.font.color.type is not None:
        dst_run.font.color.rgb = src_run.font.color.rgb


def _append_run(paragraph, text: str) -> None:
    source = paragraph.runs[-1] if paragraph.runs else None
    run = paragraph.add_run(text)
    if source is not None:
        _copy_run_format(source, run)


def _insert_paragraph_after(paragraph, text: str, template=None):
    import copy

    from docx.text.paragraph import Paragraph
    from docx.oxml import OxmlElement

    new_p = OxmlElement("w:p")
    if template is not None and template._p.pPr is not None:
        new_p.append(copy.deepcopy(template._p.pPr))
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    run = new_para.add_run(text)
    if template is not None and template.runs:
        _copy_run_format(template.runs[-1], run)
    return new_para


def _insert_paragraph_before(paragraph, text: str, template=None):
    new_para = _insert_paragraph_after(paragraph, text, template)
    paragraph._p.addprevious(new_para._p)
    return new_para
