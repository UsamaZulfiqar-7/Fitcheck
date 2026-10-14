import io

import pytest

import config
from core import cv_builder, extraction
from core.extraction import ExtractionError
from tests.sample_files import (CV_TEXT, JOB_TEXT, cv_docx_bytes, scanned_pdf_bytes, text_image_bytes, text_pdf_bytes)

needs_ocr = pytest.mark.skipif(not extraction.ocr_available(), reason="Tesseract not installed")


def test_txt_docx_pdf_extraction():
    txt = extraction.extract_text(CV_TEXT.encode(), "cv.txt", extraction.RESUME_TYPES)
    docx_ = extraction.extract_text(cv_docx_bytes(), "cv.docx", extraction.RESUME_TYPES)
    pdf = extraction.extract_text(text_pdf_bytes(), "cv.pdf", extraction.RESUME_TYPES)
    for res in (txt, docx_, pdf):
        assert "Ayesha Khan" in res.text and "Pandas" in res.text
    assert pdf.method == "embedded PDF text"


def test_docx_tables_are_read():
    import docx
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Skills", "Python, SQL, Excel and Tableau"
    buf = io.BytesIO()
    d.save(buf)
    assert "Python, SQL" in extraction.extract_text(buf.getvalue(), "t.docx", ["docx"]).text


def test_validation_errors_are_friendly():
    with pytest.raises(ExtractionError, match="empty"):
        extraction.extract_text(b"", "a.txt", ["txt"])
    with pytest.raises(ExtractionError, match="Unsupported"):
        extraction.extract_text(b"hello world", "a.exe", ["txt"])
    with pytest.raises(ExtractionError, match="not a valid PDF"):
        extraction.extract_text(b"not a pdf at all", "a.pdf", ["pdf"])
    with pytest.raises(ExtractionError, match="not a valid Word"):
        extraction.extract_text(b"not a docx", "a.docx", ["docx"])
    with pytest.raises(ExtractionError, match="Almost no text"):
        extraction.extract_text(b"hi", "a.txt", ["txt"])


def test_oversized_file(monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_MB", 0)
    with pytest.raises(ExtractionError, match="larger"):
        extraction.extract_text(b"x" * 10, "a.txt", ["txt"])


def test_corrupt_pdf_and_image():
    with pytest.raises(ExtractionError):
        extraction.extract_text(b"%PDF-1.4 garbage garbage garbage", "a.pdf", ["pdf"])
    with pytest.raises(ExtractionError, match="could not be opened"):
        extraction.extract_text(b"not an image", "a.png", ["png"])


@needs_ocr
def test_image_ocr():
    res = extraction.extract_text(text_image_bytes(), "job.png", extraction.JOB_TYPES)
    assert res.method == "OCR (image)"
    assert "SQL" in res.text and "Power BI" in res.text


@needs_ocr
def test_scanned_pdf_uses_ocr():
    res = extraction.extract_text(scanned_pdf_bytes(), "scan.pdf", extraction.JOB_TYPES)
    assert res.method.startswith("OCR") and "Python" in res.text


def test_ocr_missing_gives_install_help(monkeypatch):
    monkeypatch.setattr(extraction, "ocr_available", lambda: False)
    with pytest.raises(ExtractionError, match="Tesseract"):
        extraction.extract_text(text_image_bytes(), "job.png", extraction.JOB_TYPES)


# ---- CV builder ---------------------------------------------------------------------------------------
def test_parse_and_roundtrip_keeps_content():
    sections = cv_builder.parse_cv(CV_TEXT)
    names = cv_builder.section_names(sections)
    assert names == ["SUMMARY", "EDUCATION", "SKILLS", "PROJECTS", "EXPERIENCE"]
    text = cv_builder.cv_to_text(sections)
    for line in CV_TEXT.splitlines():
        if line.strip():
            assert line.strip() in text


def test_add_skill_to_comma_list_preserves_everything_else():
    sections = cv_builder.parse_cv(CV_TEXT)
    new = cv_builder.add_skill(sections, "SQL", "SKILLS")
    before, after = cv_builder.cv_to_text(sections), cv_builder.cv_to_text(new)
    assert "Python, Excel, Pandas, Data Cleaning, Communication, SQL" in after
    assert before.replace("Communication\n", "Communication, SQL\n") == after       # nothing else changed
    assert cv_builder.cv_to_text(sections) == before                                 # input not mutated


def test_add_skill_bullets_and_new_section():
    sections = cv_builder.parse_cv("Name\n\nPROJECTS\n- one\n- two\n")
    out = cv_builder.add_skill(sections, "Docker", "PROJECTS")
    assert cv_builder.cv_to_text(out).rstrip().endswith("- Docker")
    out2 = cv_builder.add_skill(sections, "Docker", "Technical Skills")
    assert cv_builder.section_names(out2)[-1] == "Technical Skills"
    assert cv_builder.section_preview(sections, "Technical Skills").startswith("(this section does not exist")


def test_duplicate_detection():
    sections = cv_builder.parse_cv(CV_TEXT)
    assert cv_builder.skill_already_in_cv(sections, "python")
    assert cv_builder.skill_already_in_cv(sections, "Pandas")
    assert not cv_builder.skill_already_in_cv(sections, "Power BI")
    with pytest.raises(ValueError):
        cv_builder.add_skill(sections, "  ", "SKILLS")


def test_section_options_prefers_skill_sections():
    opts = cv_builder.skills_section_options(cv_builder.parse_cv(CV_TEXT))
    assert opts[0] == "SKILLS" and "Technical Skills" in opts


def test_exports_are_valid_files():
    import docx
    sections = cv_builder.add_skill(cv_builder.parse_cv(CV_TEXT), "SQL", "SKILLS")
    data = cv_builder.to_docx(sections)
    doc = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Ayesha Khan" in text and "SQL" in text and "Sales dashboard" in text
    pdf = cv_builder.to_pdf(sections)
    assert pdf.startswith(b"%PDF")
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf)) as p:
        assert "Ayesha Khan" in (p.pages[0].extract_text() or "")


def test_docx_in_place_edit_keeps_original_formatting():
    import docx
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    h = d.add_paragraph()
    r = h.add_run("SKILLS")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    body = d.add_paragraph()
    body_run = body.add_run("Python, Excel")
    body_run.italic = True
    body_run.font.name = "Georgia"
    d.add_paragraph("EXPERIENCE")
    d.add_paragraph("Intern at Example Co")
    buf = io.BytesIO()
    d.save(buf)
    out = cv_builder.add_skill_to_docx(buf.getvalue(), "SQL", "SKILLS")
    doc = docx.Document(io.BytesIO(out))
    texts = [p.text for p in doc.paragraphs]
    assert "Python, Excel, SQL" in texts
    edited = next(p for p in doc.paragraphs if p.text == "Python, Excel, SQL")
    assert edited.runs[-1].italic and edited.runs[-1].font.name == "Georgia"      # formatting copied
    heading = next(p for p in doc.paragraphs if p.text == "SKILLS")
    assert heading.runs[0].bold and heading.runs[0].font.size == Pt(14)
    assert "Intern at Example Co" in texts and texts.count("SKILLS") == 1           # nothing else touched


def test_docx_adds_missing_section_before_experience():
    import docx
    d = docx.Document()
    d.add_paragraph("Ayesha Khan")
    d.add_paragraph("EXPERIENCE")
    d.add_paragraph("Intern at Example Co")
    buf = io.BytesIO()
    d.save(buf)
    out = cv_builder.add_skill_to_docx(buf.getvalue(), "Docker", "Technical Skills")
    texts = [p.text for p in docx.Document(io.BytesIO(out)).paragraphs]
    assert texts.index("Technical Skills") < texts.index("EXPERIENCE")
    assert texts[texts.index("Technical Skills") + 1] == "Docker"
